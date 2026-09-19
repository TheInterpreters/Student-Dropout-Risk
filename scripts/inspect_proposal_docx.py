from pathlib import Path
import sys

from docx import Document

sys.stdout.reconfigure(encoding="utf-8")


root = Path(__file__).resolve().parents[1]
path = Path(sys.argv[1]) if len(sys.argv) > 1 else root / "XDS_Project_Proposal_Draft.docx"
if not path.is_absolute():
    path = root / path
document = Document(path)

print(f"file={path}")
print(f"paragraphs={len(document.paragraphs)} tables={len(document.tables)}")
all_text = " ".join(paragraph.text for paragraph in document.paragraphs)
all_text += " " + " ".join(
    cell.text for table in document.tables for row in table.rows for cell in row.cells
)
print(f"words={len(all_text.split())}")
for section_index, section in enumerate(document.sections):
    usable_width = section.page_width - section.left_margin - section.right_margin
    print(f"section_{section_index}_usable_width_inches={usable_width / 914400:.2f}")
for index, paragraph in enumerate(document.paragraphs):
    if paragraph.text.strip():
        print(f"P{index}: [{paragraph.style.name}] {paragraph.text}")

for table_index, table in enumerate(document.tables):
    grid_widths = [
        int(column.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w"))
        for column in table._tbl.xpath("./w:tblGrid/w:gridCol")
    ]
    print(f"T{table_index}_grid_width_inches={sum(grid_widths) / 1440:.2f}")
    for row_index, row in enumerate(table.rows):
        values = [cell.text.replace("\n", " / ") for cell in row.cells]
        print(f"T{table_index}R{row_index}: {' | '.join(values)}")
