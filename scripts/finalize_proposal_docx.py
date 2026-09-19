"""Create the submission candidate from the synchronized proposal draft.

This script is idempotent: it finds content by text prefix instead of paragraph number,
refuses ambiguous matches, and writes a separate final file without overwriting the draft.
"""

from __future__ import annotations

from pathlib import Path
import re
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "XDS_Project_Proposal_Draft.docx"
OUTPUT = ROOT / "XDS_Project_Proposal_Submission_Ready.docx"


def find_one(document: Document, prefix: str):
    matches = [p for p in document.paragraphs if p.text.strip().startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph beginning {prefix!r}; found {len(matches)}")
    return matches[0]


def replace_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def strip_comment_parts(path: Path) -> None:
    """Remove dormant Google Docs/Word comment parts from the submission package."""
    excluded = {"word/comments.xml", "word/commentsExtended.xml"}

    def is_excluded(name: str) -> bool:
        return name in excluded or name.startswith("customXML/")

    temporary = path.with_suffix(".comments-stripped.docx")
    with ZipFile(path, "r") as source, ZipFile(temporary, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            if is_excluded(item.filename):
                continue
            payload = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                root = ElementTree.fromstring(payload)
                for child in list(root):
                    if is_excluded(child.attrib.get("PartName", "").lstrip("/")):
                        root.remove(child)
                payload = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            elif item.filename == "word/_rels/document.xml.rels":
                root = ElementTree.fromstring(payload)
                for child in list(root):
                    if child.attrib.get("Type", "").endswith(("/comments", "/commentsExtended", "/customXml")):
                        root.remove(child)
                payload = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(item, payload)
    temporary.replace(path)


document = Document(SOURCE)

# Reclaim vertical space toward the guideline's 3-4 page report length without touching
# font size, the workflow diagram, or wording: tighten default paragraph spacing/margins.
normal_style = next(style for style in document.styles if style.name.casefold() == "normal")
normal_style.paragraph_format.space_after = Pt(0)
normal_style.paragraph_format.line_spacing = 1.0
for section in document.sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

team_table = document.tables[0]

member_updates = {
    "1)": "1) Aye Khin Khin Hpone (Yolanda) — 125970",
    "2)": "2) Nguyen Liem Son (Lucas) — 126729",
    "3)": "3) Witchayda Theppithuk (Noey) — 127419",
    "4)": "4) Han Htoo Zaw — 127305",
    "5 )": "5) Pyae Sone Han — 127366",
}
replace_text(find_one(document, "Please write your name here"), "Group Members")
for prefix, text in member_updates.items():
    replace_text(find_one(document, prefix), text)

replace_text(
    find_one(document, "We propose to predict first-year undergraduate dropout risk"),
    "Early identification of dropout risk can enable timely academic or financial support, but an unreliable explanation can misdirect scarce advising resources or produce unjustified conclusions about a student. This project examines first-year undergraduate dropout risk using structured institutional data and TabICL, a peer-reviewed tabular foundation model (TFM) for in-context learning (Qu et al., 2025). TFMs achieve competitive results on small tabular datasets (Hollmann et al., 2025), but explanation quality is multidimensional and must be evaluated quantitatively and with users, not through anecdotal plots alone (Nauta et al., 2023). The principal contribution is therefore evidence about the reliability and practical value of local explanations. Three questions guide the study: RQ1: does TabICL add predictive value over simpler baselines; RQ2: are its local explanations faithful and stable; RQ3: do they help users anticipate model outputs and identify feasible support actions?",
)

replace_text(
    find_one(document, "The explanation is designed for a university academic advisor"),
    "The explanation is designed for a university student-support officer reviewing currently enrolled first-year students immediately after first-semester results are finalised and before second-semester support decisions. For a student flagged as high risk, the officer decides whether to prioritise that student for an outreach meeting and referral to appropriate academic or financial support. Missing a genuinely at-risk student may prevent timely contact, whereas excessive flagging consumes limited support capacity and can erode trust. The explanation must therefore clarify the model's output without presenting associations as causal diagnoses or recommendations.",
)

scope_paragraph = find_one(document, "The core project uses the UCI tabular dataset")
replace_text(
    scope_paragraph,
    "All predictive modelling and explanation evaluation use one openly licensed UCI dataset. Short, consented written responses from the human evaluation form the only unstructured component and remain linked to the study cases. No external review corpus or second modelling dataset is included, keeping the prediction task, explanations and human evidence coherently connected.",
)
for run in scope_paragraph.runs:
    run.font.highlight_color = None
paragraph_shading = scope_paragraph._p.get_or_add_pPr().find(qn("w:shd"))
if paragraph_shading is not None:
    scope_paragraph._p.get_or_add_pPr().remove(paragraph_shading)

replace_text(
    find_one(document, "Supplementary text data."),
    "Prospective qualitative data. A voluntary convenience sample from the Asian Institute of Technology (AIT) community will provide brief written rationales, suggested support actions and comments on missing information while evaluating pseudonymised UCI cases. These consented, anonymised responses form the project's only unstructured dataset, remain linked only to study case IDs, and are analysed qualitatively for comprehension, actionability and information gaps. No AIT student records will be collected, and responses will not enter the dropout model or be presented as explanations of real outcomes.",
)

replace_text(
    find_one(document, "Primary modeling data."),
    "Primary modelling data. ‘Predict Students’ Dropout and Academic Success’ (Realinho, Vieira Martins, Machado, & Baptista, 2021), UCI Machine Learning Repository dataset 697, is distributed under CC BY 4.0 and documented in a peer-reviewed data paper (Realinho et al., 2022). It contains 4,424 student records and 36 predictors spanning demographic, socio-economic, admission, administrative, academic and macroeconomic information, with labels Dropout, Enrolled and Graduate. The distributed version has no direct identifiers, and exploratory analysis found no missing values or duplicate rows.",
)

replace_text(
    find_one(document, "The primary prediction task is binary"),
    "The primary prediction task is binary: Dropout (1) versus Graduate (0). Enrolled records are excluded as unresolved outcomes rather than confirmed non-dropout cases; this simplifies interpretation but may introduce selection bias and limit applicability to similar unresolved students. A meaningful share of Dropout records (17.3%, versus 3.4% of Graduate records) show no first-semester evaluations, so some labelled dropouts may have left before the decision point rather than after it; this timing ambiguity will be reported as an explicit eligibility limitation. The resulting class distribution will be documented. Subject to time and class support, sensitivity analysis will retain Enrolled as a separate class or group it with non-dropout, while primary conclusions stay tied to the pre-specified binary task.",
)

replace_text(
    find_one(document, "The prediction time is immediately after first-semester results"),
    "The prediction time is immediately after first-semester results. Only variables available by then may enter the model: enrolment/admission records, demographic and socio-economic fields, administrative status, macroeconomic indicators, and first-semester outcomes. Second-semester curricular variables are excluded as temporal leakage. First-semester grades and approved units are predictive but already historical, so they are treated as immutable in counterfactual recourse.",
)

replace_text(
    find_one(document, "We will compare TabPFN and TabICL"),
    "A fixed, stratified 70/15/15 train/validation/test split supports identical comparisons. Before test access, protocol amendments added a native main-effects Explainable Boosting Machine (EBM) and froze all audit rules without changing the split, seed, target or feature windows. Validation ROC-AUC was 0.952 for TabICL 2.2.0 (four estimators), 0.946 for EBM, 0.945 for logistic regression and XGBoost, and 0.923 for a depth-3 tree: a modest complex-model gain, not superiority. EBM uses raw named variables, explicit nominal types and no interactions. The policy ranks each cohort and flags the top 20% with stable tie-breaking; no universal probability threshold is claimed. Among 213 validation dropouts and 108 available flags, TabICL attained the theoretical maximum 0.507 recall with 1.000 precision. At admission time, EBM (0.8753), XGBoost (0.8752) and TabICL (0.8739) were practically tied. One final same-test-set evaluation will report discrimination, capacity metrics, calibration, runtime and paired 2,000-resample bootstrap intervals for model differences.",
)

replace_text(
    find_one(document, "Because a TFM is not a tree ensemble"),
    "As a TFM is not a tree ensemble, TreeSHAP exactness will not be claimed. Model-agnostic permutation SHAP (Lundberg & Lee, 2017) explains dropout probability using 25 training-only background rows and 67 evaluations. Signed member attributions are summed within each semantic group, preserving additivity, before groups are ranked by absolute total. Deletion masks those same groups with training references, recomputes derived variables, and compares the cumulative absolute probability change with 50 random group orders; the primary statistic is ranked-minus-random area with case-level uncertainty. Similar curves indicate inadequate fidelity, and masking measures reliance, not intervention. Validation checks reconstructed TabICL probabilities within 4.1e-8 and recovered known white-box attribution ranks with Spearman 1.0. Grouped permutation importance supplies global reliance, while ALE examines three validation-selected numerical effects with less off-manifold risk than PDP under correlation.",
)

replace_text(
    find_one(document, "Following the rubric, we will pre-select 20 held-out cases"),
    "The stability evaluation deterministically selects four cases from each of five predicted-risk strata (20 total) without labels. For each case, the explainer runs 10 times, varying seed and training-background sample separately. Mean pairwise top-five overlap—shared features divided by five—is primary, with rank correlation secondary. Continuous perturbations draw Gaussian noise with standard deviation equal to 1% of the feature's training standard deviation, clip to training bounds, recompute dependencies, and qualify only when absolute probability change is at most 0.01. Mean overlap below the validation-frozen 0.80 warning level is displayed, not hidden through selective reruns (Tiukhova et al., 2024).",
)

replace_text(
    find_one(document, "We will recruit 8–10 participants where feasible"),
    "Following forward-simulation evaluation (Hase & Bansal, 2020), the human study targets 8–10 voluntary AIT participants, with a minimum of 4–5 split evenly by seeded assignment. After standardised instructions and one unscored practice case, both groups receive the same label-free panel of five flagged and five unflagged profiles in random order; the explanation group additionally sees feature contributions. Neither group receives the output, probability, capacity boundary or correctness feedback until all 10 responses are locked. The interface records predictions, confidence, time and written responses. Participant-level accuracy is primary; group accuracy and time are compared descriptively, with two or three observed comments and coded responses (Section 5.7). Unless support professionals participate, the study tests proxy-user comprehensibility, not professional effectiveness, and makes no significance claim.",
)

replace_text(
    find_one(document, "Before modeling, we will classify features"),
    "Before modelling, features are classified by temporal availability and feasibility. Age, nationality, prior qualifications and completed first-semester outcomes remain fixed. Potentially resolvable tuition or debt conditions may change only when realistic at the decision point, distinguishing institutional support from student-controlled action. For a flagged case, the search finds the smallest feasible change that moves it outside the cohort's top-20% set after recomputing its rank while all other scores stay fixed; no universal threshold is assumed, and no-recourse cases are explicit. These outputs describe model behaviour, not causal promises (Karimi et al., 2021).",
)

replace_text(
    find_one(document, "For one representative held-out student"),
    "The end-to-end workflow produces a frozen comparison of TabICL and its baselines, quantitative explanation evidence, constrained counterfactuals, subgroup diagnostics, human-evaluation findings, a completed model card and reproducible report artefacts. A lightweight browser-based prototype provides two controlled modes: a decision-support demonstration showing the frozen prediction, local explanation, reliability warnings and feasible support options, and a blinded study view that suppresses the answer until the participant responds, loading the same versioned preprocessing and model artefacts used in evaluation as a research demonstration, not a production AIT service.",
)

replace_text(
    find_one(document, "5.6  Interpretation Walkthrough (Worked Example)"),
    "5.6  System Architecture, Workflow and Expected Outputs",
)

replace_text(
    find_one(document, "After applying a pre-specified relevance filter"),
    "After the prediction task, each participant explains the judgement, proposes a support action, and flags missing information. A pre-defined coding guide groups responses into comprehension, actionability and information-gap categories; two reviewers independently code at least 20% of responses to report raw agreement, then resolve disagreements and report aggregate counts with paraphrased examples. This qualitative analysis complements the forward-simulation results but does not train, validate or explain the dropout model.",
)

replace_text(
    find_one(document, "5.7  Supplementary Text Coverage Analysis"),
    "5.7  Qualitative Analysis of Human-Study Responses",
)

# Replace the prose-only walkthrough with a numbered, editable Word workflow. The snake
# layout preserves chronological order while fitting nine substantive stages on one page.
workflow_paragraph = find_one(document, "The end-to-end workflow produces")
workflow_table = document.add_table(rows=5, cols=5)
workflow_table.alignment = WD_TABLE_ALIGNMENT.CENTER
workflow_table.autofit = False
workflow_items = {
    (0, 0): ("1. Decision frame\nintended user · decision · timing\nbinary target", "D9EAF7"),
    (0, 1): ("→", None),
    (0, 2): ("2. Data design\nUCI audit · feature windows\nfrozen 70/15/15 IDs", "D9EAF7"),
    (0, 3): ("→", None),
    (0, 4): ("3. Model development\ntrain-only preparation\nTabICL + LR/tree/XGB/EBM", "E2F0D9"),
    (1, 4): ("↓", None),
    (2, 4): ("4. Validation freeze\nmodel · capacity policy\nSHAP groups/budget", "E2F0D9"),
    (2, 3): ("←", None),
    (2, 2): ("5. One-time final test\nperformance · calibration\npaired uncertainty · runtime", "E2F0D9"),
    (2, 1): ("←", None),
    (2, 0): ("6. Local explanations\npreselected held-out cases\nmodel-agnostic SHAP", "FCE4D6"),
    (3, 0): ("↓", None),
    (4, 0): ("7. Explanation evidence\ndeletion + white-box · stability\nperturbation · recourse", "FCE4D6"),
    (4, 1): ("→", None),
    (4, 2): ("8. Decision-support study\nsubgroup/proxy audit · prototype\nblinded A/B study + coding", "E4DFEC"),
    (4, 3): ("→", None),
    (4, 4): ("9. Reporting package\nmodel card · limitations\nreproducible report + prototype", "E4DFEC"),
}
for row_index, row in enumerate(workflow_table.rows):
    for column_index, cell in enumerate(row.cells):
        label, fill = workflow_items.get((row_index, column_index), ("", None))
        is_box = fill is not None
        cell.width = Inches(1.85 if column_index % 2 == 0 else 0.25)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(label)
        run.bold = is_box
        run.font.size = Pt(7.2 if is_box else 11)
        run.font.color.rgb = RGBColor(31, 78, 121)
        properties = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            border = OxmlElement(f"w:{edge}")
            border.set(qn("w:val"), "single" if is_box else "nil")
            border.set(qn("w:sz"), "6")
            border.set(qn("w:color"), "5B9BD5")
            borders.append(border)
        properties.append(borders)
        if fill:
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), fill)
            properties.append(shading)
# Fix Word's table grid so stage columns are wide and connector columns stay narrow.
workflow_widths = [1.85, 0.25, 1.85, 0.25, 1.85]
for grid_column, width in zip(
    workflow_table._tbl.xpath("./w:tblGrid/w:gridCol"), workflow_widths, strict=True
):
    grid_column.set(qn("w:w"), str(int(width * 1440)))
for row in workflow_table.rows:
    row_properties = row._tr.get_or_add_trPr()
    row_properties.append(OxmlElement("w:cantSplit"))
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True
workflow_paragraph._p.addnext(workflow_table._tbl)
caption = document.add_paragraph(
    "Figure 1. Proposed research workflow. Model, capacity policy and explanation protocols are frozen on validation data before one-time final testing. Ethics, governance and reproducibility apply throughout."
)
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.space_after = Pt(6)
for run in caption.runs:
    run.italic = True
    run.font.size = Pt(8)
workflow_table._tbl.addnext(caption._p)

replace_text(
    find_one(document, "The UCI dataset is public"),
    "The UCI dataset is publicly available under CC BY 4.0 with no direct identifiers; its citation and licence will be preserved, the raw file unchanged, and all transformations documented. AIT human-study participation will be voluntary, informed-consent based, and penalty-free to decline or withdraw. No names, AIT identifiers or administrative records will be collected; raw free text stays outside version control, accidental identifiers are removed, and only aggregate or paraphrased findings are reported. Before recruitment, the team will confirm with the course instructor whether institutional approval is required, obtaining it where applicable.",
)

replace_text(
    find_one(document, "We will explicitly list protected"),
    "Protected attributes—gender, nationality, age, international status, displaced status and special needs—are explicit. Every group count is retained; performance metrics are suppressed rather than silently omitted below 30 records, and unavailable comparisons are stated. Eligible groups receive flag rate, true- and false-positive rates, precision, mean risk and calibration gap under the exact capacity labels. All non-protected model features are screened as potential proxies using Cramér's V for categorical features and correlation-ratio eta for continuous features. These diagnostics are descriptive, not proof of fairness.",
)

replace_text(
    find_one(document, "We will also examine plausible proxies"),
    "Plausible proxies will also be examined: parental occupation and qualification may encode socio-economic position, course and application mode may reflect demographic sorting, and debtor, tuition-fee and scholarship variables may encode financial disadvantage. Removing a protected attribute does not eliminate correlated proxy information, so influential proxies and limitations will be documented even where they weaken the model's apparent suitability. The system supports outreach prioritisation, not autonomous denial of education or services.",
)

replace_text(
    find_one(document, "The main model is a tabular foundation model"),
    "TabICL combines learned representations with in-context information from a reference set and fails simulatability and decomposability: its reasoning cannot be recovered from a coefficient list or small tree, motivating Section 5.2's post-hoc protocol. By contrast, the main-effects EBM is intrinsically interpretable: each prediction is exactly the logistic transform of an intercept plus one learned shape-function contribution per feature. Its validation probabilities reconstruct from those terms within 3.4e-16, and global term strengths plus fixed low-, median- and high-risk local decompositions are exported. Logistic regression and the depth-3 tree provide further simulatable comparisons. TabICL's validation ROC-AUC exceeds EBM by only 0.006, so the project does not presume that complexity is operationally justified; final same-test-set performance, calibration, runtime and explanation evidence will determine whether TabICL remains a research comparison or EBM is the more defensible deployment candidate.",
)

replace_text(
    find_one(document, "The contribution is not another accuracy-only application"),
    "The proposed contribution extends beyond accuracy, asking whether TabICL's post-hoc local explanations merit trust and whether an intrinsically interpretable EBM offers a more defensible alternative. Same-test-set baselines, deletion faithfulness, white-box recovery, repeated-run and perturbation stability, human forward simulation, qualitative evidence and feasibility-constrained counterfactuals form a coherent evaluation. Limitations include the single-institution UCI context, exclusion of unresolved Enrolled outcomes, approximate model-agnostic attribution, non-causal observational data, a small convenience sample, and likely non-professional audience proxies; therefore, the AIT study tests comprehensibility, not predictive validity or deployment suitability at AIT. If recruitment reaches only the minimum sample, the study stays descriptive. Findings remain valuable if EBM matches or exceeds TabICL, explanations are unstable or unfaithful, participants gain no benefit, or feasible recourse is unavailable.",
)

replace_text(
    find_one(document, "The repository separates untouched source data"),
    "The repository separates source data, EDA, modelling, explanation evaluation, study materials and outputs. A pinned environment supports a fresh clone. Seed 42 and source-row IDs define the split; protocol 1.2, 26 in-scope tests, four baselines, capacity-consistent recourse, exact EBM decomposition, TabICL/SHAP integration, grouped importance, ALE, explicit subgroup suppression and proxy tables are implemented. Before final submission, one command will regenerate every report artefact and another will launch the prototype from frozen artefacts. The model card records intended use, data, final metrics, explanation evidence, subgroup results and limitations.",
)

replace_text(find_one(document, "1. Overview"), "1. Introduction, Aim and Research Questions")
replace_text(find_one(document, "2. Target Audience and Decision"), "2. Decision Context and Intended User")
replace_text(find_one(document, "3. Dataset"), "3. Data Sources")
replace_text(find_one(document, "4. Problem Formulation"), "4. Prediction Task and Study Design")
replace_text(find_one(document, "5. Method"), "5. Proposed Methodology")
replace_text(find_one(document, "5.1  Model and Baseline"), "5.1  Predictive Modelling and Evaluation")
replace_text(find_one(document, "5.2  Faithfulness Evidence"), "5.2  Explanation Method and Faithfulness Evaluation")
replace_text(find_one(document, "5.3  Stability Evidence"), "5.3  Explanation Stability Evaluation")
replace_text(find_one(document, "5.4  Human Evidence"), "5.4  Human Evaluation")
replace_text(find_one(document, "5.5  Actionability"), "5.5  Actionability and Counterfactual Analysis")
replace_text(find_one(document, "6. Ethics and Data Handling"), "6. Ethics, Fairness and Data Governance")
replace_text(find_one(document, "7. Model Complexity Justification"), "7. Feasibility and Model Complexity")
replace_text(find_one(document, "8. Reproducibility Plan"), "8. Reproducibility and Documentation")
replace_text(find_one(document, "9. Team and Task Division"), "9. Work Plan and Responsibilities")
replace_text(
    find_one(document, "10. Contribution"),
    "9. Expected Contribution, Limitations and Contingencies",
)

replace_text(
    find_one(document, "The audience, prediction time, target, split"),
    "The audience, prediction time, binary target, leakage-safe feature window and evaluation protocol have been agreed jointly. Each lead will produce reproducible artefacts, while a second member will review the corresponding code and claims before submission:",
)

references_heading = find_one(document, "References")
replace_text(references_heading, "Appendix A. References")
references_heading.paragraph_format.page_break_before = True
availability_paragraph = document.add_paragraph(
    "Code and data availability. The versioned project repository will be released at https://github.com/TheInterpreters/Proposal_EDA with code, tests, non-sensitive tables and reproduction commands; raw human free text will remain excluded."
)
for run in availability_paragraph.runs:
    run.font.size = Pt(10)
references_heading._p.addprevious(availability_paragraph._p)

references = find_one(document, "Realinho, V.")
reference_text = "\n".join(
    (
        "Realinho, V., Vieira Martins, M., Machado, J., & Baptista, L. (2021). Predict Students’ Dropout and Academic Success [Dataset]. UCI Machine Learning Repository. https://doi.org/10.24432/C5MC89",
        "Realinho, V., Machado, J., Baptista, L., & Martins, M. V. (2022). Predicting student dropout and academic success. Data, 7(11), 146. https://doi.org/10.3390/data7110146",
        "Hollmann, N., Müller, S., Purucker, L., Krishnakumar, A., Körfer, M., Hoo, S. B., Schirrmeister, R. T., & Hutter, F. (2025). Accurate predictions on small data with a tabular foundation model. Nature, 637, 319–326. https://doi.org/10.1038/s41586-024-08328-6",
        "Qu, J., Holzmüller, D., Varoquaux, G., & Le Morvan, M. (2025). TabICL: A tabular foundation model for in-context learning on large data. Proceedings of the 42nd International Conference on Machine Learning, PMLR 267, 50817–50847. https://proceedings.mlr.press/v267/qu25d.html",
        "Lou, Y., Caruana, R., Gehrke, J., & Hooker, G. (2013). Accurate intelligible models with pairwise interactions. Proceedings of the 19th ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, 623–631. https://doi.org/10.1145/2487575.2487579",
        "Lundberg, S. M., & Lee, S.-I. (2017). A unified approach to interpreting model predictions. Advances in Neural Information Processing Systems, 30, 4765–4774. https://proceedings.neurips.cc/paper/2017/hash/8a20a8621978632d76c43dfd28b67767-Abstract.html",
        "Nauta, M., Trienes, J., Pathak, S., Nguyen, E., Peters, M., Schmitt, Y., Schlötterer, J., van Keulen, M., & Seifert, C. (2023). From anecdotal evidence to quantitative evaluation methods: A systematic review on evaluating explainable AI. ACM Computing Surveys, 55(13s), Article 295. https://doi.org/10.1145/3583558",
        "Hase, P., & Bansal, M. (2020). Evaluating explainable AI: Which algorithmic explanations help users predict model behavior? Proceedings of the 58th Annual Meeting of the Association for Computational Linguistics, 5540–5552. https://doi.org/10.18653/v1/2020.acl-main.491",
        "Karimi, A.-H., Schölkopf, B., & Valera, I. (2021). Algorithmic recourse: From counterfactual explanations to interventions. Proceedings of the 2021 ACM Conference on Fairness, Accountability, and Transparency, 353–362. https://doi.org/10.1145/3442188.3445899",
        "Tiukhova, E., Vemuri, P., López Flores, N., Islind, A. S., Óskarsdóttir, M., Poelmans, S., Baesens, B., & Snoeck, M. (2024). Explainable learning analytics: Assessing the stability of student success prediction models by means of explainable AI. Decision Support Systems, 182, 114229. https://doi.org/10.1016/j.dss.2024.114229",
    )
)
replace_text(references, reference_text)
references.alignment = WD_ALIGN_PARAGRAPH.LEFT

# The proposal identifies all group members on the title page. A task-allocation table
# and timeline are intentionally omitted because neither is required and both duplicate
# the methodological workflow within a limited page budget.
for removable_prefix in (
    "9. Work Plan and Responsibilities",
    "The audience, prediction time, binary target",
):
    paragraph = find_one(document, removable_prefix)
    paragraph._element.getparent().remove(paragraph._element)
team_table._element.getparent().remove(team_table._element)

# The source draft contains an old Google Docs comment/highlight around the scope
# paragraph. Submission copies should not display editing markup or yellow highlighting.
for paragraph in document.paragraphs:
    for highlight in list(paragraph._p.xpath(".//w:highlight")):
        highlight.getparent().remove(highlight)
    for tag in ("w:commentRangeStart", "w:commentRangeEnd", "w:commentReference"):
        for marker in list(paragraph._p.xpath(f".//{tag}")):
            marker.getparent().remove(marker)

document.save(OUTPUT)
strip_comment_parts(OUTPUT)

check = Document(OUTPUT)
all_text = "\n".join(p.text for p in check.paragraphs)
table_text = "\n".join(
    cell.text for table in check.tables for row in table.rows for cell in row.cells
)
assert "70/15/15 train/validation/test" in all_text
assert "semantic group" in all_text
assert "TabICL: A tabular foundation model" in all_text
assert "TabICLv2" not in all_text
assert "arXiv" not in all_text
assert "selection bias" in all_text
assert "theoretical maximum 0.507 recall" in all_text
assert "At admission time" in all_text
assert "Explainable Boosting Machine" in all_text
assert "3.4e-16" in all_text
assert "TabICL + LR/tree/XGB/EBM" in table_text
assert "protocol 1.2" in all_text
assert "9. Expected Contribution, Limitations and Contingencies" in all_text
assert "No AIT student records will be collected" in all_text
assert "AIT study tests comprehensibility" in all_text
assert "Three questions" in all_text
assert "Qualitative Analysis of Human-Study Responses" in all_text
assert "No external review corpus" in all_text
assert "shared features divided by five" in all_text
assert "fresh clone" in all_text
assert "https://github.com/TheInterpreters/Proposal_EDA" in all_text
assert "Ethics, governance and reproducibility apply throughout" in all_text
assert "1. Decision frame" in table_text
assert "9. Reporting package" in table_text
assert "blinded A/B study" in table_text
assert "(Criterion" not in all_text
assert "rubric" not in all_text.lower()
assert "Jaccard" not in all_text
assert "final.csv" not in all_text
assert "UK university" not in all_text
prose_text = "\n".join(p.text for p in check.paragraphs if not p.text.startswith("Realinho, V."))
first_person = re.findall(r"\bI\b|\b(?:me|my|mine|we|us|our|ours)\b", prose_text)
assert not first_person, f"First-person pronouns remain: {first_person}"
assert not re.search(r"(?:^|[.!?]\s+)Because\b", prose_text), "A sentence starts with Because"
assert "Work Plan and Responsibilities" not in all_text
assert len(check.tables) == 1
assert "advisor" not in (all_text + table_text).lower()
assert "timeline" not in (all_text + table_text).lower()
print(f"Created: {OUTPUT}")
